#!/usr/bin/env python3
"""
Read mapped outputs from a recalculated RealAI pro forma workbook.

Requires the workbook to have been recalculated first (use the xlsx skill's
scripts/recalc.py). Reads cached calculated values via openpyxl(data_only=True).

Usage:
    python read_outputs.py --workbook path/to/proforma.xlsx --manifest path/to/manifest.md

Outputs JSON to stdout:
    {
        "outputs": {
            "year1_cap_rate": {"cell": "Assumptions!H6", "value": 0.067},
            ...
        },
        "errors": [
            {"key": "...", "cell": "...", "error": "#DIV/0!"}
        ],
        "status": "success" | "errors_found"
    }
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

try:
    import yaml
except ImportError:
    print("ERROR: PyYAML not installed. Run: pip install pyyaml --break-system-packages", file=sys.stderr)
    sys.exit(5)

try:
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
except ImportError:
    print("ERROR: openpyxl not installed.", file=sys.stderr)
    sys.exit(5)


_CELL_RE = re.compile(r"^(?P<sheet>[^!]+)!(?P<col>[A-Z]+)(?P<row>\d+)$")
_EXCEL_ERRORS = {"#REF!", "#DIV/0!", "#VALUE!", "#N/A", "#NAME?", "#NUM!", "#NULL!"}


def load_manifest(path: Path) -> dict:
    text = path.read_text(encoding="utf-8")
    m = re.search(r"```ya?ml\s*\n(.*?)\n```", text, re.DOTALL)
    if not m:
        print("ERROR: Manifest has no fenced YAML block.", file=sys.stderr)
        sys.exit(2)
    return yaml.safe_load(m.group(1))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--workbook", required=True, type=Path)
    ap.add_argument("--manifest", required=True, type=Path)
    args = ap.parse_args()

    if not args.workbook.exists():
        print(f"ERROR: Workbook not found: {args.workbook}", file=sys.stderr)
        sys.exit(5)

    manifest = load_manifest(args.manifest)
    wb = load_workbook(args.workbook, data_only=True)

    outputs = {}
    errors = []

    for key, odef in manifest.get("outputs", {}).items():
        cell_ref = odef["cell"]
        m = _CELL_RE.match(cell_ref)
        if not m:
            errors.append({"key": key, "cell": cell_ref, "error": "bad_cell_ref"})
            continue
        sheet = m.group("sheet")
        col_idx = column_index_from_string(m.group("col"))
        row = int(m.group("row"))
        try:
            value = wb[sheet].cell(row=row, column=col_idx).value
        except KeyError:
            errors.append({"key": key, "cell": cell_ref, "error": f"sheet {sheet!r} not found"})
            continue
        if isinstance(value, str) and value in _EXCEL_ERRORS:
            errors.append({"key": key, "cell": cell_ref, "error": value})
            continue
        outputs[key] = {"cell": cell_ref, "value": value, "semantic_role": odef.get("semantic_role")}

    result = {
        "status": "errors_found" if errors else "success",
        "outputs": outputs,
        "errors": errors,
    }
    print(json.dumps(result, indent=2, default=str))


if __name__ == "__main__":
    main()
