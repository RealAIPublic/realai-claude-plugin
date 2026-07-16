"""RealAI docx applicator — port of the original docx-formatter agent.

Re-skins a finished .docx to the RealAI visual language without altering text
runs, table values, or chart data. Verifies byte-identity of text content
before returning. On mismatch, discards the output.

Usage:
    python apply_docx_brand.py --in /path/in.docx --out /path/out.docx
"""

from __future__ import annotations

import argparse
import shutil
import sys
from pathlib import Path
from typing import List

try:
    from docx import Document
    from docx.shared import Pt, Twips, RGBColor
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
except ImportError:
    print("python-docx is required. Install via: pip install python-docx --break-system-packages")
    sys.exit(2)


# ---- Token values (mirror /skills/realai-brand/tokens.md). ----

PALETTE = {
    "content_primary":   "151513",
    "content_tertiary":  "6B6A65",
    "accent":            "F4633A",
    "canvas":            "FFFFFF",
    "greige":            "FAF9F6",
    "blackish":          "22262A",
    "hairline":          "E4E3DF",
    "slate":             "5B6B7A",
    "gold":              "B8923B",
    "status_red":        "C0392B",
    "status_yellow":     "C9A227",
}

FONT_SERIF = "Ivar Headline"
FONT_DISPLAY = "Ivar Display"
FONT_SANS = "Source Sans 3"
FALLBACK_SERIF = "Georgia"
FALLBACK_SANS = "Calibri"

TYPE_SCALE = {
    "title":      28,
    "h1":         20,
    "h2":         15,
    "h3":         12,
    "eyebrow":     9,
    "body":       11,
    "body_small": 10,
    "caption":     9,
    "disclaimer":  9,
    "table_cell": 10,
    "table_hdr":   9,
}


# ---- Verification spine: snapshot text runs in order, ignoring whitespace-only. ----

def snapshot_text_runs(doc: Document) -> List[str]:
    """Ordered list of every visible text run in the document.

    Includes paragraphs in body, in tables (cell by cell), and in headers/footers.
    Whitespace-only runs are filtered out.
    """
    runs: List[str] = []

    def walk_paragraphs(paragraphs):
        for p in paragraphs:
            for r in p.runs:
                t = r.text
                if t and t.strip():
                    runs.append(t)

    walk_paragraphs(doc.paragraphs)
    for tbl in doc.tables:
        for row in tbl.rows:
            for cell in row.cells:
                walk_paragraphs(cell.paragraphs)
                for inner in cell.tables:
                    for r2 in inner.rows:
                        for c2 in r2.cells:
                            walk_paragraphs(c2.paragraphs)

    for sec in doc.sections:
        walk_paragraphs(sec.header.paragraphs)
        walk_paragraphs(sec.footer.paragraphs)

    return runs


def assert_text_identity(before: List[str], after: List[str]) -> None:
    if before == after:
        return
    raise RuntimeError(
        "Text-run verification FAILED — formatting altered content. "
        f"Input had {len(before)} non-empty runs; output had {len(after)}. "
        "Output discarded per the verification spine."
    )


# ---- Style layer application. ----

def _set_run_props(run, font_name: str | None = None, size_pt: int | None = None, color_hex: str | None = None):
    """Strip conflicting direct formatting, keep semantic bold/italic/underline."""
    rPr = run._element.get_or_add_rPr()

    # Strip highlight, font color shadings.
    for tag in ("w:highlight", "w:shd"):
        el = rPr.find(qn(tag))
        if el is not None:
            rPr.remove(el)

    if font_name:
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rFonts.set(qn(attr), font_name)
    if size_pt is not None:
        sz = rPr.find(qn("w:sz"))
        if sz is None:
            sz = OxmlElement("w:sz")
            rPr.append(sz)
        sz.set(qn("w:val"), str(size_pt * 2))  # half-points
    if color_hex is not None:
        col = rPr.find(qn("w:color"))
        if col is None:
            col = OxmlElement("w:color")
            rPr.append(col)
        col.set(qn("w:val"), color_hex)


def _apply_paragraph_spacing(paragraph, before_pt: int, after_pt: int, line_pt: int | None = None, auto_line: bool = False):
    pPr = paragraph._element.get_or_add_pPr()
    spc = pPr.find(qn("w:spacing"))
    if spc is None:
        spc = OxmlElement("w:spacing")
        pPr.append(spc)
    spc.set(qn("w:before"), str(before_pt * 20))  # twips
    spc.set(qn("w:after"), str(after_pt * 20))
    if line_pt is not None:
        spc.set(qn("w:line"), str(line_pt))
        spc.set(qn("w:lineRule"), "auto" if auto_line else "exact")


def _detect_section_title(paragraph) -> str | None:
    """Return 'h2', 'h3', or None per the section-title detection rule.

    Treat as H2 when the paragraph is short (<=8 words), on its own line,
    and visually distinguished (bold, non-body color, ALL CAPS, or already a heading).
    Short bold labels ending in ':' become H3 (no rule).
    """
    text = paragraph.text.strip()
    if not text:
        return None
    words = text.split()
    if len(words) > 8:
        return None

    is_caps = text.isupper() and any(c.isalpha() for c in text)
    is_bold = any(r.bold for r in paragraph.runs if r.bold)
    style_name = (paragraph.style.name or "").lower()
    looks_like_heading = style_name.startswith("heading") or style_name in {"title", "subtitle"}

    if not (is_bold or is_caps or looks_like_heading):
        return None

    if is_bold and text.endswith(":"):
        return "h3"
    return "h2"


def _add_h2_rule(paragraph):
    pPr = paragraph._element.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = pBdr.find(qn("w:bottom"))
    if bottom is None:
        bottom = OxmlElement("w:bottom")
        pBdr.append(bottom)
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")         # 0.75 pt in eighths-of-a-point
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), PALETTE["accent"])


def apply_style_layer(doc: Document) -> None:
    """Restyle the document to RealAI tokens. Text runs untouched."""

    # 1. Walk paragraphs: strip conflicting run-level direct formatting and apply
    #    body styles by default; promote detected section titles.
    for p in doc.paragraphs:
        promotion = _detect_section_title(p)

        if promotion == "h2":
            _apply_paragraph_spacing(p, before_pt=14, after_pt=4)
            for r in p.runs:
                _set_run_props(r, font_name=FONT_SERIF, size_pt=TYPE_SCALE["h2"], color_hex=PALETTE["content_primary"])
            _add_h2_rule(p)
        elif promotion == "h3":
            _apply_paragraph_spacing(p, before_pt=10, after_pt=3)
            for r in p.runs:
                _set_run_props(r, font_name=FONT_SERIF, size_pt=TYPE_SCALE["h3"], color_hex=PALETTE["content_primary"])
        else:
            # Body
            _apply_paragraph_spacing(p, before_pt=0, after_pt=7, line_pt=276, auto_line=True)
            for r in p.runs:
                _set_run_props(r, font_name=FONT_SANS, size_pt=TYPE_SCALE["body"], color_hex=PALETTE["content_primary"])

    # 2. Tables: header fill greige, hairline borders, header bottom rule accent.
    for tbl in doc.tables:
        for row_idx, row in enumerate(tbl.rows):
            for cell in row.cells:
                tcPr = cell._element.get_or_add_tcPr()
                if row_idx == 0:
                    # Header fill
                    shd = tcPr.find(qn("w:shd"))
                    if shd is None:
                        shd = OxmlElement("w:shd")
                        tcPr.append(shd)
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:fill"), PALETTE["greige"])
                # Borders
                tcBorders = tcPr.find(qn("w:tcBorders"))
                if tcBorders is None:
                    tcBorders = OxmlElement("w:tcBorders")
                    tcPr.append(tcBorders)
                for side in ("top", "bottom", "left", "right"):
                    b = tcBorders.find(qn(f"w:{side}"))
                    if b is None:
                        b = OxmlElement(f"w:{side}")
                        tcBorders.append(b)
                    is_header_bottom = (row_idx == 0 and side == "bottom")
                    b.set(qn("w:val"), "single")
                    b.set(qn("w:sz"), "12" if is_header_bottom else "4")
                    b.set(qn("w:color"), PALETTE["accent"] if is_header_bottom else PALETTE["hairline"])
                # Cell text style
                for p in cell.paragraphs:
                    for r in p.runs:
                        _set_run_props(
                            r,
                            font_name=FONT_SANS,
                            size_pt=TYPE_SCALE["table_hdr" if row_idx == 0 else "table_cell"],
                            color_hex=PALETTE["content_primary"],
                        )


# ---- Entry point. ----

def main(argv=None):
    parser = argparse.ArgumentParser(description="Apply RealAI house brand to a .docx.")
    parser.add_argument("--in", dest="src", required=True)
    parser.add_argument("--out", dest="dst", required=True)
    args = parser.parse_args(argv)

    src = Path(args.src)
    dst = Path(args.dst)
    if not src.exists():
        print(f"Input not found: {src}", file=sys.stderr)
        return 2

    work = dst.with_suffix(dst.suffix + ".work")
    shutil.copyfile(src, work)
    doc = Document(work)

    before = snapshot_text_runs(doc)
    apply_style_layer(doc)
    after = snapshot_text_runs(doc)

    try:
        assert_text_identity(before, after)
    except RuntimeError as e:
        work.unlink(missing_ok=True)
        print(f"FAILED: {e}", file=sys.stderr)
        return 3

    doc.save(work)
    shutil.move(work, dst)
    print(f"OK: {dst}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
