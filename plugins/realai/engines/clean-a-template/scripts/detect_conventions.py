#!/usr/bin/env python3
"""Detect workbook conventions: unit scale, time axes, signs, and iterative calc."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

UNIT_MARKERS = [
    (re.compile(r"\$\s*000|\$\s*in\s*000|dollars?\s*in\s*thousands|in\s*thousands", re.I), "$000"),
    (re.compile(r"\$\s*mm|\$\s*m\b|in\s*millions|dollars?\s*in\s*millions", re.I), "$M"),
    (re.compile(r"whole\s*dollars|actual\s*dollars|in\s*dollars", re.I), "$"),
]
ANCHOR_LABELS = re.compile(
    r"(offer\s*price|purchase\s*price|acquisition\s*price|contract\s*price|total\s*capitalization|loan\s*amount|debt\s*proceeds|total\s*equity|total\s*sources|total\s*uses|stabilized\s*noi|year\s*1\s*noi)",
    re.I,
)
EXPENSE_LABELS = re.compile(r"(tax|insurance|utilities|repairs?|maintenance|payroll|admin|management\s*fee|operating\s*expense|opex)", re.I)
MONTH_TEXT = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|sept|oct|nov|dec)\b", re.I)
YEAR_TEXT = re.compile(r"\b(20\d{2}|19\d{2}|year\s*[1-9][0-9]?)\b", re.I)


def nearest_label(ws, row: int, col: int) -> str | None:
    for dc in range(1, 5):
        c = col - dc
        if c < 1:
            break
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v and not v.startswith("="):
            return v.strip()
    for dr in range(1, 3):
        r = row - dr
        if r < 1:
            break
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v and not v.startswith("="):
            return v.strip()
    return None


def marker_scan(wb) -> list[dict]:
    hits = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    for pattern, scale in UNIT_MARKERS:
                        if pattern.search(cell.value):
                            hits.append({"sheet": ws.title, "cell": cell.coordinate, "scale": scale, "text": cell.value[:120]})
    return hits


def anchor_amounts(wb) -> list[dict]:
    anchors = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    label = nearest_label(ws, cell.row, cell.column)
                    if label and ANCHOR_LABELS.search(label):
                        v = float(cell.value)
                        abs_v = abs(v)
                        if abs_v >= 1_000_000:
                            inferred = "$"
                        elif 1_000 <= abs_v < 1_000_000:
                            inferred = "$000"
                        elif 10 <= abs_v < 1_000:
                            inferred = "$M"
                        else:
                            inferred = "unknown"
                        anchors.append(
                            {
                                "sheet": ws.title,
                                "cell": cell.coordinate,
                                "label": label[:120],
                                "value_class": "redacted_numeric_anchor",
                                "inferred_scale": inferred,
                            }
                        )
    return anchors


def detect_unit_scale(wb) -> dict:
    markers = marker_scan(wb)
    if markers:
        counts = Counter(m["scale"] for m in markers)
        scale, count = counts.most_common(1)[0]
        return {"scale": scale, "confidence": "high", "source": "explicit_marker", "evidence": markers[:20]}
    anchors = anchor_amounts(wb)
    usable = [a["inferred_scale"] for a in anchors if a["inferred_scale"] != "unknown"]
    if usable:
        scale, count = Counter(usable).most_common(1)[0]
        confidence = "high" if count >= 2 else "medium"
        return {"scale": scale, "confidence": confidence, "source": "anchor_amount_magnitude", "evidence": anchors[:20]}
    return {"scale": "unknown", "confidence": "low", "source": "not_detected", "evidence": []}


def period_kind(values: list[Any]) -> str | None:
    if len(values) < 3:
        return None
    text = " ".join(str(v) for v in values if v is not None)
    if MONTH_TEXT.search(text):
        return "monthly"
    if sum(1 for v in values if isinstance(v, (datetime, date))) >= 3:
        # Use count rather than spacing: Excel models often store dates with formulas.
        return "monthly_or_date_based"
    if len(YEAR_TEXT.findall(text)) >= 3:
        return "annual"
    nums = [v for v in values if isinstance(v, (int, float)) and not isinstance(v, bool)]
    if len(nums) >= 3 and all(1900 <= float(v) <= 2100 for v in nums[: min(len(nums), 10)]):
        return "annual"
    return None


def detect_time_axes(wb) -> list[dict]:
    candidates = []
    for ws in wb.worksheets:
        for row_idx in range(1, (ws.max_row or 0) + 1):
            row_vals = [ws.cell(row=row_idx, column=c).value for c in range(1, (ws.max_column or 0) + 1)]
            nonblank_positions = [(i + 1, v) for i, v in enumerate(row_vals) if v not in (None, "")]
            if len(nonblank_positions) < 3:
                continue
            kind = period_kind([v for _, v in nonblank_positions])
            if kind:
                candidates.append(
                    {
                        "sheet": ws.title,
                        "row": row_idx,
                        "granularity": "monthly" if kind == "monthly_or_date_based" and len(nonblank_positions) >= 12 else kind,
                        "period_count": len(nonblank_positions),
                        "first_column": nonblank_positions[0][0],
                        "last_column": nonblank_positions[-1][0],
                        "role_guess": "historical_source_data" if re.search(r"t-?12|trailing", ws.title, re.I) else "projection_or_output",
                    }
                )
    # Prefer compact output.
    candidates.sort(key=lambda x: (0 if x["role_guess"] == "projection_or_output" else 1, -x["period_count"]))
    return candidates[:30]


def detect_expense_sign(wb_values) -> dict:
    signs = Counter()
    evidence = []
    for ws in wb_values.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    label = nearest_label(ws, cell.row, cell.column)
                    if label and EXPENSE_LABELS.search(label):
                        sign = "negative" if cell.value < 0 else "positive"
                        signs[sign] += 1
                        if len(evidence) < 20:
                            evidence.append({"sheet": ws.title, "cell": cell.coordinate, "label": label[:120], "sign": sign})
    if not signs:
        return {"expense_sign": "unknown", "confidence": "low", "evidence": []}
    sign, count = signs.most_common(1)[0]
    confidence = "high" if count >= 5 else "medium"
    return {"expense_sign": sign, "confidence": confidence, "evidence": evidence}


def iterative_calc(wb) -> dict:
    calc = getattr(wb, "calculation", None) or getattr(wb, "calculation_properties", None)
    iterate = bool(getattr(calc, "iterate", False)) if calc else False
    return {"iterative_calc_enabled": iterate}


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_file")
    parser.add_argument("--out")
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input_file, data_only=False, read_only=False, keep_links=True)
    wb_values = openpyxl.load_workbook(args.input_file, data_only=True, read_only=False, keep_links=True)
    result = {
        "file": args.input_file,
        "unit_scale": detect_unit_scale(wb),
        "time_axis_candidates": detect_time_axes(wb),
        "expense_sign": detect_expense_sign(wb_values),
        "iterative_calc": iterative_calc(wb),
        "note": "If multiple time axes exist, choose the primary underwriting projection axis in Phase 1 judgment log.",
    }
    text = json.dumps(result, indent=2, default=str)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# END detect_conventions.py
