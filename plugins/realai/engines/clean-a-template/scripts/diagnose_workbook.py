#!/usr/bin/env python3
"""Workbook diagnostics: errors, external refs, named ranges, hidden data, and formula smells."""

from __future__ import annotations

import argparse
import json
import re
import zipfile
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any

import openpyxl

ERROR_VALUES = {"#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!", "#GETTING_DATA"}
EXTERNAL_RE = re.compile(r"\[[^\]]+\.(?:xlsx|xlsm|xlsb|xls)\]", re.I)
CELL_REF_RE = re.compile(r"(?:'[^']+'|[A-Za-z_][A-Za-z0-9_ .&-]*!)?\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?")
NUMERIC_LITERAL_RE = re.compile(r"(?<![A-Z0-9])[-+]?\d+(?:\.\d+)?%?(?![A-Z0-9])", re.I)
TODO_RE = re.compile(r"\b(todo|fixme|xxx|check this|plug|placeholder)\b", re.I)


def formula_signature(formula: str) -> str:
    s = CELL_REF_RE.sub("REF", formula or "")
    s = re.sub(r"\s+", "", s).upper()
    return s


def count_errors(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=False, keep_links=True)
    by_sheet = Counter()
    by_type = Counter()
    samples = []
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value in ERROR_VALUES:
                    by_sheet[ws.title] += 1
                    by_type[cell.value] += 1
                    if len(samples) < 50:
                        samples.append({"sheet": ws.title, "cell": cell.coordinate, "error": cell.value})
    return {"total": sum(by_type.values()), "by_sheet": dict(by_sheet), "by_type": dict(by_type), "samples": samples}


def scan_formulas(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    external_refs = []
    formula_patterns = []
    magic_numbers = []
    todos = []
    for ws in wb.worksheets:
        for row_idx in range(1, (ws.max_row or 0) + 1):
            row_formulas = []
            for col_idx in range(1, (ws.max_column or 0) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                v = cell.value
                if isinstance(v, str):
                    if TODO_RE.search(v):
                        todos.append({"sheet": ws.title, "cell": cell.coordinate, "text": v[:160]})
                    if v.startswith("="):
                        if EXTERNAL_RE.search(v):
                            external_refs.append({"sheet": ws.title, "cell": cell.coordinate, "formula": v[:240]})
                        row_formulas.append((cell.coordinate, formula_signature(v), v))
                        nums = [n for n in NUMERIC_LITERAL_RE.findall(v) if n not in {"0", "1", "-1", "12", "4", "365", "100"}]
                        if nums and len(magic_numbers) < 100:
                            magic_numbers.append({"sheet": ws.title, "cell": cell.coordinate, "numbers": nums[:8], "formula": v[:180]})
            if len(row_formulas) >= 4:
                counts = Counter(sig for _, sig, _ in row_formulas)
                if len(counts) > 1:
                    main_sig, main_count = counts.most_common(1)[0]
                    outliers = [coord for coord, sig, _ in row_formulas if sig != main_sig]
                    if main_count >= 3 and len(outliers) <= max(3, len(row_formulas) // 4):
                        formula_patterns.append({"sheet": ws.title, "row": row_idx, "outlier_cells": outliers[:20], "pattern_count": len(counts)})
    return {
        "external_formula_refs": external_refs,
        "formula_pattern_warnings": formula_patterns[:100],
        "hardcoded_formula_constants": magic_numbers,
        "todo_markers": todos[:100],
    }


def broken_names(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    out = []
    for name, defined_name in wb.defined_names.items():
        try:
            destinations = list(defined_name.destinations)
            if not destinations:
                continue
            for sheet, ref in destinations:
                if sheet not in wb.sheetnames:
                    out.append({"name": name, "issue": "missing_sheet", "target": f"{sheet}!{ref}"})
        except Exception as exc:  # noqa: BLE001
            out.append({"name": name, "issue": "cannot_resolve", "error": str(exc)})
    return out


def package_findings(path: Path) -> dict:
    findings = {"external_parts": [], "vba_project": False, "custom_xml_parts": []}
    try:
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            findings["vba_project"] = any(n.lower() == "xl/vbaproject.bin" for n in names)
            findings["external_parts"] = [n for n in names if n.lower().startswith("xl/externallinks/") or n.lower().startswith("xl/connections")]
            findings["custom_xml_parts"] = [n for n in names if n.lower().startswith("customxml/")][:50]
    except zipfile.BadZipFile:
        pass
    return findings


def hidden_sheets(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=False, read_only=False, keep_links=True)
    return [{"sheet": ws.title, "state": ws.sheet_state, "max_row": ws.max_row, "max_column": ws.max_column} for ws in wb.worksheets if ws.sheet_state != "visible"]


def issue_list(result: dict) -> list[dict]:
    issues = []
    if result["errors"]["total"]:
        issues.append({"id": "C.1", "severity": "critical", "issue": "Formula/value errors present", "count": result["errors"]["total"]})
    if result["formula_scan"]["external_formula_refs"] or result["package"]["external_parts"]:
        issues.append({"id": "C.3", "severity": "critical", "issue": "External references or connection parts present"})
    if result["broken_named_ranges"]:
        issues.append({"id": "C.4", "severity": "critical", "issue": "Broken or unresolved named ranges", "count": len(result["broken_named_ranges"])})
    if result["package"]["vba_project"]:
        issues.append({"id": "C.5", "severity": "critical", "issue": "VBA project present; .xlsm/macros unsupported in v1"})
    for key, issue_id, label in [
        ("formula_pattern_warnings", "W.1", "Inconsistent formula patterns"),
        ("hardcoded_formula_constants", "W.2", "Hardcoded constants inside formulas"),
        ("todo_markers", "W.8", "TODO/FIXME markers"),
    ]:
        count = len(result["formula_scan"].get(key, []))
        if count:
            issues.append({"id": issue_id, "severity": "warning", "issue": label, "count": count})
    if result["hidden_sheets"]:
        issues.append({"id": "I.2", "severity": "info", "issue": "Hidden sheets present", "count": len(result["hidden_sheets"])})
    return issues


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_file")
    parser.add_argument("--out")
    args = parser.parse_args()
    path = Path(args.input_file)
    result = {
        "file": str(path),
        "errors": count_errors(path),
        "formula_scan": scan_formulas(path),
        "broken_named_ranges": broken_names(path),
        "hidden_sheets": hidden_sheets(path),
        "package": package_findings(path),
    }
    result["issues"] = issue_list(result)
    text = json.dumps(result, indent=2, default=str)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# END diagnose_workbook.py
