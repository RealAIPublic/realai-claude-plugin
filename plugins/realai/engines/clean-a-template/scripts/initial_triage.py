#!/usr/bin/env python3
"""Initial workbook triage for template preparation.

Outputs JSON with file type, workbook size, unsupported features, external links,
and the formula-reevaluation handoff. This script does not modify the workbook.
"""

from __future__ import annotations

import argparse
import json
import sys
import zipfile
from pathlib import Path

import openpyxl


def zip_parts(path: Path) -> list[str]:
    try:
        with zipfile.ZipFile(path) as zf:
            return zf.namelist()
    except zipfile.BadZipFile:
        return []


def find_external_parts(parts: list[str]) -> list[str]:
    markers = []
    for name in parts:
        lower = name.lower()
        if lower.startswith("xl/externallinks/"):
            markers.append(name)
        if lower in {"xl/connections.xml", "xl/querytables.xml"}:
            markers.append(name)
        if lower.startswith("xl/querytables/") or lower.startswith("xl/connections/"):
            markers.append(name)
    return sorted(set(markers))


def workbook_counts(path: Path) -> dict:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=False, keep_links=True)
    sheets = []
    total_cells = 0
    total_visible_cells = 0
    for ws in wb.worksheets:
        cells = int((ws.max_row or 0) * (ws.max_column or 0))
        total_cells += cells
        if ws.sheet_state == "visible":
            total_visible_cells += cells
        sheets.append(
            {
                "name": ws.title,
                "state": ws.sheet_state,
                "max_row": ws.max_row,
                "max_column": ws.max_column,
                "cell_area": cells,
            }
        )
    return {
        "sheet_count": len(wb.worksheets),
        "visible_sheet_count": sum(1 for s in sheets if s["state"] == "visible"),
        "total_cell_area": total_cells,
        "visible_cell_area": total_visible_cells,
        "sheets": sheets,
        "defined_name_count": len(list(wb.defined_names.items())) if hasattr(wb.defined_names, "items") else 0,
        "properties": {
            "creator_present": bool(getattr(wb.properties, "creator", None)),
            "last_modified_by_present": bool(getattr(wb.properties, "lastModifiedBy", None)),
            "title_present": bool(getattr(wb.properties, "title", None)),
            "subject_present": bool(getattr(wb.properties, "subject", None)),
        },
    }


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_file")
    parser.add_argument("--out")
    args = parser.parse_args()

    path = Path(args.input_file)
    result: dict = {
        "file": str(path),
        "extension": path.suffix.lower(),
        "supported_extension": path.suffix.lower() == ".xlsx",
        "issues": [],
        "refusal_reasons": [],
        "formula_recalc": {
            "method": "excel_skill_scripts_recalc_py",
            "command": "cd <excel_skill_root> && python scripts/recalc.py <absolute_workbook> [timeout_seconds]",
            "note": "Phase 3 must run sniff_test.py, which invokes the Excel skill recalc.py after the sniff payload is written.",
        },
    }

    parts = zip_parts(path)
    result["zip_parts_checked"] = bool(parts)
    result["has_vba_project"] = "xl/vbaproject.bin" in {p.lower() for p in parts}
    result["external_connection_parts"] = find_external_parts(parts)

    if path.suffix.lower() == ".xlsm" or result["has_vba_project"]:
        # v1.3: macro-enabled workbooks are no longer flat-refused. They route
        # through the xlsm gate, which analyzes the VBA project (never executes
        # it), scans it for deal-data leakage, and emits a macro-stripped .xlsx
        # when the proof obligations hold. The pipeline then continues on the
        # converted file.
        result["requires_xlsm_gate"] = True
        result["xlsm_gate"] = {
            "script": "tp_scripts/xlsm_triage_convert.py",
            "command": ("python tp_scripts/xlsm_triage_convert.py <model.xlsm> "
                        "--out xlsm_report.json --convert-out model_converted.xlsx "
                        "[--decisions-json decisions.json]"),
            "note": ("Run the gate BEFORE any other Phase 1 script. verdict=strippable: "
                     "continue the pipeline on the converted .xlsx and carry "
                     "removed_macro_inventory into _PreparationAudit. verdict=needs_review: "
                     "present the Function/event inventory to the user; re-run with "
                     "--approve-strip after approval. verdict=not_strippable: refuse with "
                     "the report's reason (formulas depend on VBA)."),
        }
    elif path.suffix.lower() != ".xlsx":
        result["refusal_reasons"].append("Only macro-free .xlsx workbooks are supported in v1.")

    try:
        counts = workbook_counts(path)
        result.update(counts)
        if counts["visible_sheet_count"] > 25 and counts["visible_cell_area"] > 500_000:
            result["refusal_reasons"].append(
                f"Workbook has {counts['visible_sheet_count']} visible sheets and {counts['visible_cell_area']} visible cell area; this exceeds the safe automated-mapping threshold."
            )
    except Exception as exc:  # noqa: BLE001
        result["refusal_reasons"].append(f"Workbook could not be opened cleanly: {type(exc).__name__}: {exc}")

    if result["external_connection_parts"]:
        result["issues"].append(
            {
                "severity": "critical_review",
                "issue": "External data connection parts found",
                "parts": result["external_connection_parts"],
                "recommended_action": "Proceed only if the user approves static conversion or confirms the links are not required.",
            }
        )

    if result.get("requires_xlsm_gate") and not result["refusal_reasons"]:
        result["status"] = "xlsm_gate_required"
    else:
        result["status"] = "not_compatible" if result["refusal_reasons"] else "triage_passed"

    text = json.dumps(result, indent=2, default=str)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# END initial_triage.py
