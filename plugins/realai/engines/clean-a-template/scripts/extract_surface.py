#!/usr/bin/env python3
"""Extract sheet roles, regions, candidate scalar inputs, output reads, and table schemas."""

from __future__ import annotations

import argparse
import json
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

DATA_TABLE_NAMES = re.compile(r"(t-?12|trailing|rent\s*roll|unit\s*mix|comps|budget|draw|import)", re.I)
INPUT_SURFACE_NAMES = re.compile(r"(assumptions?|inputs?|pro\s*forma|proforma|key\s*terms?|deal\s*overview)", re.I)
OUTPUT_NAMES = re.compile(r"(returns?|summary|waterfall|one\s*pager|dashboard|metrics)", re.I)
DEAD_NAMES = re.compile(r"^(sheet\d+|old|scratch|backup|copy\s*of|v\d+)", re.I)

OUTPUT_VOCAB = re.compile(
    r"(irr|equity\s*multiple|moic|noi|net\s*operating\s*income|cap\s*rate|dscr|debt\s*yield|cash\s*on\s*cash|exit\s*value|sale\s*proceeds|yield\s*on\s*cost)",
    re.I,
)

INPUT_LABEL_HINTS = re.compile(
    r"(price|cost|rent|growth|cap\s*rate|rate|ltv|loan|debt|vacancy|occupancy|expense|tax|insurance|payroll|management|fee|units?|sq\.?\s*ft|sf|hold|term|amort|exit|noi|capex|renovation|premium|preferred|promote|split|hurdle)",
    re.I,
)

FORMULA_SHEET_REF = re.compile(r"(?:'([^']+)'|([A-Za-z_][A-Za-z0-9_ .&-]*))!")


def classify_value(value: Any) -> str:
    if value is None:
        return "blank"
    if isinstance(value, str) and value.startswith("="):
        return "formula"
    if isinstance(value, str):
        return "text"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, (int, float)):
        return "number"
    if isinstance(value, (datetime, date)):
        return "date"
    return "other"


def value_class(value: Any) -> str:
    cls = classify_value(value)
    if cls == "number":
        v = abs(float(value))
        if v >= 1_000_000:
            return "large_number"
        if v >= 1_000:
            return "medium_number"
        if 0 <= v <= 1:
            return "rate_or_ratio"
    if cls == "text":
        return "text"
    return cls


def census(ws) -> Counter:
    counts: Counter = Counter()
    for row in ws.iter_rows():
        for cell in row:
            counts[classify_value(cell.value)] += 1
    return counts


def referenced_sheets(wb) -> set[str]:
    out: set[str] = set()
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    for m in FORMULA_SHEET_REF.finditer(cell.value):
                        out.add((m.group(1) or m.group(2) or "").strip())
    return out


def detect_regions(ws, row_step: int = 8) -> list[dict]:
    regions: list[dict] = []
    current: dict | None = None
    for start in range(1, (ws.max_row or 0) + 1, row_step):
        end = min(start + row_step - 1, ws.max_row or start)
        counts = Counter()
        for row in ws.iter_rows(min_row=start, max_row=end):
            for cell in row:
                counts[classify_value(cell.value)] += 1
        nonblank = sum(counts[k] for k in ["formula", "number", "text", "date", "boolean", "other"])
        if nonblank < 4:
            role = "empty"
        else:
            formula_ratio = counts["formula"] / nonblank
            numeric_ratio = (counts["number"] + counts["date"] + counts["boolean"]) / nonblank
            if formula_ratio > 0.7:
                role = "calculation"
            elif numeric_ratio > 0.25 and counts["formula"] < counts["number"] + counts["text"]:
                role = "input_like"
            else:
                role = "mixed"
        if current and current["role"] == role:
            current["rows"][1] = end
        else:
            if current:
                regions.append(current)
            current = {"rows": [start, end], "role": role}
    if current:
        regions.append(current)
    return regions


def classify_sheet(ws, counts: Counter, referenced_by_others: bool) -> tuple[str, str, list[str]]:
    name = ws.title
    nonblank = sum(counts[k] for k in counts if k != "blank")
    formula_ratio = counts["formula"] / nonblank if nonblank else 0
    number_ratio = counts["number"] / nonblank if nonblank else 0
    reasons: list[str] = []
    if nonblank == 0:
        return "dead", "high", ["empty sheet"]
    if DATA_TABLE_NAMES.search(name):
        return "data_table", "high", ["sheet name indicates data table"]
    if DEAD_NAMES.search(name) and not referenced_by_others:
        return "dead", "medium", ["sheet name suggests legacy/scratch and no inbound references"]
    if INPUT_SURFACE_NAMES.search(name):
        reasons.append("sheet name indicates input/pro forma surface")
        if formula_ratio > 0.35:
            return "mixed", "high", reasons + ["contains both formulas and hardcoded values"]
        return "input_surface", "high", reasons
    if OUTPUT_NAMES.search(name) and formula_ratio > 0.25:
        return "output", "medium", ["sheet name indicates output/summary"]
    if (ws.max_row or 0) > 80 and number_ratio > 0.45 and formula_ratio < 0.25:
        return "data_table", "medium", ["large numeric table signature"]
    if formula_ratio > 0.7:
        return "calculation", "medium", ["high formula density"]
    if formula_ratio > 0.2 and number_ratio > 0.1:
        return "mixed", "medium", ["mixed formula and numeric density"]
    return "support", "low", ["fallback classification"]


def nearest_label(ws, row: int, col: int, left: int = 4, up: int = 2) -> dict | None:
    labels: list[tuple[int, str, str, str]] = []
    for dc in range(1, left + 1):
        c = col - dc
        if c < 1:
            break
        v = ws.cell(row=row, column=c).value
        if isinstance(v, str) and v and not v.startswith("="):
            labels.append((dc, v.strip(), f"{get_column_letter(c)}{row}", "left"))
            break
    for dr in range(1, up + 1):
        r = row - dr
        if r < 1:
            break
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v and not v.startswith("="):
            labels.append((dr + 2, v.strip(), f"{get_column_letter(col)}{r}", "above"))
            break
    if not labels:
        return None
    labels.sort(key=lambda x: x[0])
    _, text, coord, direction = labels[0]
    return {"text": text[:200], "cell": coord, "direction": direction}


def in_any_table(ws, coord: str) -> bool:
    for table in ws.tables.values():
        if coord in ws[table.ref]:
            return True
    return False


def table_schemas(ws) -> list[dict]:
    out = []
    for name, table in ws.tables.items():
        out.append({"name": name, "sheet": ws.title, "ref": table.ref, "displayName": table.displayName})
    return out


def extract_candidates(ws, sheet_role: str, max_candidates: int) -> tuple[list[dict], list[dict]]:
    inputs: list[dict] = []
    outputs: list[dict] = []
    skip_scalar = sheet_role == "data_table"
    for row in ws.iter_rows():
        for cell in row:
            cls = classify_value(cell.value)
            if cls == "blank":
                continue
            label = nearest_label(ws, cell.row, cell.column)
            if label and cls in {"number", "text", "date", "boolean"}:
                if OUTPUT_VOCAB.search(label["text"]):
                    outputs.append(
                        {
                            "cell": f"{ws.title}!{cell.coordinate}",
                            "label": label,
                            "value_class": value_class(cell.value),
                            "source": "adjacent_label",
                        }
                    )
                elif not skip_scalar and INPUT_LABEL_HINTS.search(label["text"]):
                    inputs.append(
                        {
                            "cell": f"{ws.title}!{cell.coordinate}",
                            "label": label,
                            "value_class": value_class(cell.value),
                            "in_excel_table": in_any_table(ws, cell.coordinate),
                            "source": "adjacent_label",
                        }
                    )
        if len(inputs) >= max_candidates and len(outputs) >= max_candidates:
            break
    return inputs[:max_candidates], outputs[:max_candidates]


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("input_file")
    parser.add_argument("--out")
    parser.add_argument("--max-candidates-per-sheet", type=int, default=250)
    args = parser.parse_args()

    wb = openpyxl.load_workbook(args.input_file, data_only=False, read_only=False, keep_links=True)
    refs = referenced_sheets(wb)
    result = {"file": args.input_file, "sheets": [], "candidate_inputs": [], "candidate_outputs": [], "tables": []}
    for ws in wb.worksheets:
        counts = census(ws)
        role, confidence, reasons = classify_sheet(ws, counts, ws.title in refs)
        regions = detect_regions(ws) if role in {"mixed", "input_surface", "support"} else []
        sheet_info = {
            "name": ws.title,
            "state": ws.sheet_state,
            "role": role,
            "confidence": confidence,
            "reasons": reasons,
            "census": dict(counts),
            "regions": regions,
        }
        result["sheets"].append(sheet_info)
        result["tables"].extend(table_schemas(ws))
        if ws.sheet_state == "visible":
            inputs, outputs = extract_candidates(ws, role, args.max_candidates_per_sheet)
            result["candidate_inputs"].extend(inputs)
            result["candidate_outputs"].extend(outputs)
    text = json.dumps(result, indent=2, default=str)
    if args.out:
        Path(args.out).write_text(text, encoding="utf-8")
    print(text)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())

# END extract_surface.py
