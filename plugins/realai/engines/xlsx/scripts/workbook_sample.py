#!/usr/bin/env python3
"""
Create lightweight samples from spreadsheets for large-file triage.

Usage:
    python skills/xlsx/scripts/workbook_sample.py <file>
        [--rows-per-sheet N]      default 20
        [--max-cols N]            default 80
        [--output-dir DIR]        default output/samples
        [--mode values|formulas|both]   default both (xlsx/xlsm only)
        [--strategy auto|head|stratified]   default auto

Modes (xlsx/xlsm only; delimited files always use values):
    values    - computed values only
    formulas  - formula text where present, raw value otherwise
    both      - emits both per sheet (default; preferred, agents should compare)

Strategies:
    auto       - stratified when the sheet is much larger than the budget, else head
    head       - first N non-empty rows
    stratified - head + evenly spaced middle + tail

Outputs:
    - Per-sheet CSV sample files (up to 2 per sheet when mode=both)
    - sample_manifest.json with per-sheet paths and metadata
"""

import argparse
import csv
import json
import re
import sys
from collections import deque
from pathlib import Path

from openpyxl import load_workbook


def sanitize_name(name):
    return re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("_") or "sheet"


def choose_strategy(strategy_arg, max_row, rows_per_sheet):
    if strategy_arg != "auto":
        return strategy_arg
    if max_row > rows_per_sheet * 5:
        return "stratified"
    return "head"


def _write_sections(writer, head_rows, mid_rows, tail_rows):
    total = 0
    for r in head_rows:
        writer.writerow(["" if v is None else v for v in r])
        total += 1
    if mid_rows:
        writer.writerow(["--- sampled middle rows ---"])
        for r in mid_rows:
            writer.writerow(["" if v is None else v for v in r])
            total += 1
    if tail_rows:
        writer.writerow(["--- sampled tail rows ---"])
        for r in tail_rows:
            writer.writerow(["" if v is None else v for v in r])
            total += 1
    return total


def _collect_rows_worksheet(ws, rows_per_sheet, max_cols, strategy):
    """Return (head_rows, mid_rows, tail_rows) for a worksheet."""
    if strategy == "head":
        head_rows = []
        for row in ws.iter_rows(values_only=True):
            bounded = list(row[:max_cols]) if row else []
            if any(c is not None and str(c).strip() != "" for c in bounded):
                head_rows.append(bounded)
            if len(head_rows) >= rows_per_sheet:
                break
        return head_rows, [], []

    # stratified
    total_rows = ws.max_row or 0
    if total_rows <= rows_per_sheet:
        return _collect_rows_worksheet(ws, rows_per_sheet, max_cols, "head")

    head_budget = max(1, rows_per_sheet // 2)
    tail_budget = max(1, rows_per_sheet // 4)
    mid_budget = max(0, rows_per_sheet - head_budget - tail_budget)

    mid_region_start = head_budget + 1
    mid_region_end = total_rows - tail_budget
    mid_indices = set()
    if mid_budget > 0 and mid_region_end > mid_region_start:
        span = mid_region_end - mid_region_start + 1
        step = max(1, span // mid_budget)
        for k in range(mid_budget):
            idx = mid_region_start + k * step
            if idx <= mid_region_end:
                mid_indices.add(idx)

    tail_start_row = total_rows - tail_budget + 1
    head_rows = []
    mid_rows = []
    tail_buf = deque(maxlen=tail_budget)

    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        bounded = list(row[:max_cols]) if row else []
        if not any(c is not None and str(c).strip() != "" for c in bounded):
            continue
        if len(head_rows) < head_budget:
            head_rows.append(bounded)
        if i in mid_indices:
            mid_rows.append(bounded)
        if i >= tail_start_row:
            tail_buf.append(bounded)

    return head_rows, mid_rows, list(tail_buf)


def sample_workbook_one_mode(
    path, data_only, mode_label, output_dir, rows_per_sheet, max_cols, strategy_arg
):
    wb = load_workbook(path, read_only=True, data_only=data_only)
    sheet_results = []
    for i, ws in enumerate(wb.worksheets, start=1):
        strategy = choose_strategy(strategy_arg, ws.max_row or 0, rows_per_sheet)
        head_rows, mid_rows, tail_rows = _collect_rows_worksheet(
            ws, rows_per_sheet, max_cols, strategy
        )
        # Sheet-index prefix avoids collisions when sanitized names overlap.
        sampled_path = output_dir / (
            f"{i:02d}_{sanitize_name(ws.title)}.{mode_label}.sample.csv"
        )
        with open(sampled_path, "w", encoding="utf-8", newline="") as out_fh:
            writer = csv.writer(out_fh)
            rows_written = _write_sections(writer, head_rows, mid_rows, tail_rows)

        sheet_results.append({
            "sheet_name": ws.title,
            "sheet_index": i,
            "mode": mode_label,
            "strategy": strategy,
            "sample_file": str(sampled_path),
            "rows_written": rows_written,
            "source_max_row": ws.max_row or 0,
            "source_max_col": ws.max_column or 0,
        })
    wb.close()
    return sheet_results


def merge_sheet_results(results_by_mode):
    """Merge per-mode results into unified per-sheet entries."""
    per_sheet = {}
    for mode, results in results_by_mode.items():
        for r in results:
            key = (r["sheet_index"], r["sheet_name"])
            entry = per_sheet.get(key)
            if entry is None:
                entry = {
                    "sheet_name": r["sheet_name"],
                    "sheet_index": r["sheet_index"],
                    "source_max_row": r["source_max_row"],
                    "source_max_col": r["source_max_col"],
                    "strategy": r["strategy"],
                    "sample_files": {},
                    "rows_written": {},
                }
                per_sheet[key] = entry
            entry["sample_files"][mode] = r["sample_file"]
            entry["rows_written"][mode] = r["rows_written"]

    return [per_sheet[k] for k in sorted(per_sheet.keys())]


def sample_delimited(path, delimiter, rows_per_sheet, max_cols, output_dir, strategy_arg):
    sampled_path = output_dir / f"01_{sanitize_name(path.stem)}.values.sample.csv"

    # Stratified/auto needs a row count first; plain head doesn't.
    total_rows = 0
    if strategy_arg in ("auto", "stratified"):
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as fh:
            for _ in csv.reader(fh, delimiter=delimiter):
                total_rows += 1
        strategy = choose_strategy(strategy_arg, total_rows, rows_per_sheet)
    else:
        strategy = "head"

    max_seen_cols = 0
    rows_written = 0

    if strategy == "head":
        with open(path, "r", encoding="utf-8", errors="replace", newline="") as in_fh, open(
            sampled_path, "w", encoding="utf-8", newline=""
        ) as out_fh:
            reader = csv.reader(in_fh, delimiter=delimiter)
            writer = csv.writer(out_fh)
            for row in reader:
                if len(row) > max_seen_cols:
                    max_seen_cols = len(row)
                bounded = row[:max_cols]
                if any(c.strip() != "" for c in bounded):
                    writer.writerow(bounded)
                    rows_written += 1
                if rows_written >= rows_per_sheet:
                    break
    else:
        head_budget = max(1, rows_per_sheet // 2)
        tail_budget = max(1, rows_per_sheet // 4)
        mid_budget = max(0, rows_per_sheet - head_budget - tail_budget)
        tail_start = total_rows - tail_budget + 1

        mid_region_start = head_budget + 1
        mid_region_end = total_rows - tail_budget
        mid_indices = set()
        if mid_budget > 0 and mid_region_end > mid_region_start:
            span = mid_region_end - mid_region_start + 1
            step = max(1, span // mid_budget)
            for k in range(mid_budget):
                idx = mid_region_start + k * step
                if idx <= mid_region_end:
                    mid_indices.add(idx)

        head_rows = []
        mid_rows = []
        tail_buf = deque(maxlen=tail_budget)

        with open(path, "r", encoding="utf-8", errors="replace", newline="") as in_fh:
            reader = csv.reader(in_fh, delimiter=delimiter)
            for i, row in enumerate(reader, start=1):
                if len(row) > max_seen_cols:
                    max_seen_cols = len(row)
                bounded = row[:max_cols]
                if not any(c.strip() != "" for c in bounded):
                    continue
                if len(head_rows) < head_budget:
                    head_rows.append(bounded)
                if i in mid_indices:
                    mid_rows.append(bounded)
                if i >= tail_start:
                    tail_buf.append(bounded)

        with open(sampled_path, "w", encoding="utf-8", newline="") as out_fh:
            writer = csv.writer(out_fh)
            rows_written = _write_sections(writer, head_rows, mid_rows, list(tail_buf))

    entry = {
        "sheet_name": path.name,
        "sheet_index": 1,
        "strategy": strategy,
        "sample_files": {"values": str(sampled_path)},
        "rows_written": {"values": rows_written},
        "source_max_row": total_rows if total_rows > 0 else None,
        "source_max_col": max_seen_cols,
    }
    return [entry]


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("file")
    parser.add_argument("--rows-per-sheet", type=int, default=20)
    parser.add_argument("--max-cols", type=int, default=80)
    parser.add_argument("--output-dir", default="output/samples")
    parser.add_argument(
        "--mode",
        choices=["values", "formulas", "both"],
        default="both",
        help="xlsx/xlsm only; delimited files always use values",
    )
    parser.add_argument(
        "--strategy",
        choices=["auto", "head", "stratified"],
        default="auto",
        help="auto = stratified for big sheets, head otherwise",
    )
    args = parser.parse_args()

    path = Path(args.file).expanduser()
    if not path.exists():
        print(json.dumps({"status": "error", "error": f"File does not exist: {path}"}))
        sys.exit(1)

    output_dir = Path(args.output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    suffix = path.suffix.lower()
    if suffix == ".xls":
        print(json.dumps({
            "status": "error",
            "error": ".xls is not supported by openpyxl. Convert to .xlsx first."
        }))
        sys.exit(1)

    try:
        if suffix in (".xlsx", ".xlsm"):
            modes_to_run = {
                "values": [("values", True)],
                "formulas": [("formulas", False)],
                "both": [("values", True), ("formulas", False)],
            }[args.mode]

            results_by_mode = {}
            for mode_label, data_only in modes_to_run:
                results_by_mode[mode_label] = sample_workbook_one_mode(
                    path=path,
                    data_only=data_only,
                    mode_label=mode_label,
                    output_dir=output_dir,
                    rows_per_sheet=args.rows_per_sheet,
                    max_cols=args.max_cols,
                    strategy_arg=args.strategy,
                )
            samples = merge_sheet_results(results_by_mode)

        elif suffix == ".csv":
            samples = sample_delimited(
                path, ",", args.rows_per_sheet, args.max_cols, output_dir, args.strategy
            )
        elif suffix == ".tsv":
            samples = sample_delimited(
                path, "\t", args.rows_per_sheet, args.max_cols, output_dir, args.strategy
            )
        else:
            print(json.dumps({
                "status": "error",
                "error": f"Unsupported file type: {suffix}",
            }))
            sys.exit(1)

        manifest = {
            "status": "success",
            "source_file": str(path),
            "mode": args.mode,
            "strategy": args.strategy,
            "rows_per_sheet": args.rows_per_sheet,
            "max_cols": args.max_cols,
            "output_dir": str(output_dir),
            "samples": samples,
        }

        manifest_path = output_dir / "sample_manifest.json"
        manifest_path.write_text(
            json.dumps(manifest, indent=2, default=str), encoding="utf-8"
        )
        manifest["manifest_file"] = str(manifest_path)

        print(json.dumps(manifest, indent=2, default=str))
    except Exception as exc:
        print(json.dumps({"status": "error", "error": str(exc)}))
        sys.exit(1)


if __name__ == "__main__":
    main()
